import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

import pandas as pd
from openpyxl import load_workbook

from processors.normalizer import normalize_dataframe
from processors.rule_engine import classify, load_rules
from app.pipeline import process_files, export_result
from exporters.snowedge_exporter import build_snowedge_package


class TestCore(unittest.TestCase):
    def test_mapping_and_rules(self):
        raw = pd.DataFrame([
            {"IP地址": "10.0.0.1", "端口": "445/tcp", "漏洞名称": "SMB test", "风险等级": "High", "CVSS": "8.1"},
            {"IP地址": "10.0.0.2", "端口": "22", "漏洞名称": "SSH弱口令", "风险等级": "中危", "用户名": "root", "密码": "root123"},
        ])
        norm, _ = normalize_dataframe(raw, "demo.csv", "CSV")
        out = classify(norm, load_rules(ROOT / "rules"))
        self.assertEqual(out.iloc[0]["asset_ip"], "10.0.0.1")
        self.assertEqual(out.iloc[0]["severity"], "高危")
        self.assertEqual(out.iloc[0]["is_high_risk_port"], "是")
        self.assertEqual(out.iloc[0]["is_high_vulnerability"], "是")
        self.assertEqual(out.iloc[1]["is_weak_password"], "是")

    def test_manual_mapping_override(self):
        raw = pd.DataFrame([{"机器地址": "10.1.2.3", "问题": "Demo issue", "分级": "High"}])
        norm, meta = normalize_dataframe(
            raw,
            "demo.xlsx",
            "Sheet1",
            mapping_override={"机器地址": "asset_ip", "问题": "vuln_name", "分级": "severity"},
        )
        self.assertEqual(norm.iloc[0]["asset_ip"], "10.1.2.3")
        self.assertEqual(norm.iloc[0]["vuln_name"], "Demo issue")
        self.assertEqual(norm.iloc[0]["severity"], "高危")
        self.assertEqual(meta["mapping"]["机器地址"], "asset_ip")

    def test_customer_two_highs_and_weak_workbook_auto_conversion(self):
        with tempfile.TemporaryDirectory() as td:
            source = Path(td) / "客户两高一弱清单.xlsx"
            sheets = {
                "高危漏洞": [
                    ["某客户两高一弱检查结果", "", "", ""],
                    ["统计时间：2026-09-14", "", "", ""],
                    ["漏洞地址", "漏洞名称", "危险等级", "漏洞描述"],
                    ["https://demo.example:8443/login", "后台认证绕过", "高", "可绕过登录校验"],
                ],
                "高危端口": [
                    ["高危端口明细", "", ""],
                    ["IP地址", "开放端口", "端口服务"],
                    ["10.20.30.40", "6379/tcp", "redis"],
                ],
                "弱口令": [
                    ["弱口令核查结果", "", "", ""],
                    ["设备IP", "登录账号", "登录密码", "验证结果"],
                    ["10.20.30.50", "admin", "CustomerValue!", "登录验证成功"],
                ],
            }
            with pd.ExcelWriter(source, engine="openpyxl") as writer:
                for name, rows in sheets.items():
                    pd.DataFrame(rows).to_excel(writer, sheet_name=name, index=False, header=False)

            result, metadata = process_files([str(source)], str(ROOT / "rules"))
            self.assertEqual(len(result), 3)
            self.assertEqual({m["source_category"] for m in metadata}, {"高危漏洞", "高危端口", "弱口令"})
            vuln = result[result["source_category"] == "高危漏洞"].iloc[0]
            self.assertEqual(vuln["target_url"], "https://demo.example:8443/login")
            self.assertEqual(vuln["hostname"], "demo.example")
            self.assertEqual(vuln["port"], "8443")
            self.assertEqual(vuln["is_high_vulnerability"], "是")
            port = result[result["source_category"] == "高危端口"].iloc[0]
            self.assertEqual(port["is_high_risk_port"], "是")
            weak = result[result["source_category"] == "弱口令"].iloc[0]
            self.assertEqual(weak["is_weak_password"], "是")
            self.assertIn("客户原表分类=弱口令", weak["match_reason"])

    def test_password_presence_is_not_automatically_a_weak_password(self):
        raw = pd.DataFrame([{
            "IP地址": "10.1.2.3", "用户名": "analyst",
            "密码": "A-Long-Random-Phrase-2026!", "漏洞名称": "凭据记录",
        }])
        norm, _ = normalize_dataframe(raw, "demo.csv", "CSV")
        out = classify(norm, load_rules(ROOT / "rules"))
        self.assertEqual(out.iloc[0]["is_weak_password"], "否")
        self.assertEqual(out.iloc[0]["credential_status"], "发现凭据（待核验）")

    def test_end_to_end_and_branding(self):
        sample = ROOT / "samples" / "sample_vulns.csv"
        df, _ = process_files([str(sample)], str(ROOT / "rules"))
        self.assertGreater(len(df), 0)
        with tempfile.TemporaryDirectory() as td:
            out = Path(td) / "out.xlsx"
            export_result(df, str(out), [str(sample)], mask_passwords=True)
            self.assertTrue(out.exists())
            wb = load_workbook(out, read_only=True, data_only=False)
            self.assertEqual(wb.sheetnames[:6], ["汇总", "高危漏洞", "高危端口", "弱口令", "未识别数据", "全部标准化数据"])
            ws = wb["汇总"]
            values = [ws.cell(row=i, column=2).value for i in range(1, 5)]
            self.assertIn("SnowRelay", values)
            self.assertNotIn("SnowPeak", values)
            wb.close()

    def test_excel_formula_is_neutralized_and_snowedge_package_has_no_password(self):
        df, _ = process_files([str(ROOT / "samples" / "sample_vulns.csv")], str(ROOT / "rules"))
        df.loc[df.index[0], "description"] = "=HYPERLINK(\"https://invalid.example\")"
        df.loc[df.index[0], "password"] = "do-not-export"
        with tempfile.TemporaryDirectory() as td:
            out = Path(td) / "safe.xlsx"
            export_result(df, str(out), [str(ROOT / "samples" / "sample_vulns.csv")], mask_passwords=True)
            wb = load_workbook(out, read_only=True, data_only=False)
            ws = wb["全部标准化数据"]
            headers = [cell.value for cell in ws[1]]
            description_cell = ws.cell(2, headers.index("漏洞描述") + 1)
            password_cell = ws.cell(2, headers.index("密码/口令") + 1)
            self.assertEqual(description_cell.data_type, "s")
            self.assertTrue(str(description_cell.value).startswith("'="))
            self.assertEqual(password_cell.value, "••••••••")
            wb.close()

        package = build_snowedge_package(df, [str(ROOT / "samples" / "sample_vulns.csv")])
        self.assertEqual(package["schema"], "snowedge-import/1")
        self.assertFalse(package["sensitive_fields"]["password_included"])
        self.assertTrue(all("password" not in row for row in package["records"]))
        self.assertNotIn("do-not-export", str(package))


if __name__ == "__main__":
    unittest.main()
