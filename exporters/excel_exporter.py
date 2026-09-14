from __future__ import annotations

import os
import tempfile
from pathlib import Path
from typing import List

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from models.schema import EXPORT_LABELS, STANDARD_COLUMNS

PRODUCT_NAME = "SnowRelay"
PRODUCT_DESC = "Security Finding Normalization"


def _mask_password(value: object) -> str:
    s = str(value or "")
    if not s:
        return ""
    return "••••••••"


def _safe_excel_text(value: object) -> object:
    """Prevent imported text from becoming an active spreadsheet formula."""
    if not isinstance(value, str):
        return value
    if value.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def _neutralize_formulas(df: pd.DataFrame) -> pd.DataFrame:
    out = df.copy()
    for column in out.columns:
        if out[column].dtype == "object":
            out[column] = out[column].map(_safe_excel_text)
    return out


def _for_export(df: pd.DataFrame, mask_passwords: bool) -> pd.DataFrame:
    cols = [c for c in STANDARD_COLUMNS if c in df.columns]
    out = df[cols].copy()
    if mask_passwords and "password" in out.columns:
        out["password"] = out["password"].map(_mask_password)
    return _neutralize_formulas(out.rename(columns=EXPORT_LABELS))


def build_summary(df: pd.DataFrame, input_files: List[str]) -> pd.DataFrame:
    ips = {x for x in df.get("asset_ip", pd.Series(dtype=str)).astype(str) if x}
    risk_assets = {
        str(row.asset_ip) for row in df.itertuples()
        if getattr(row, "risk_type", "") and getattr(row, "asset_ip", "")
    }
    items = [
        ("产品", PRODUCT_NAME),
        ("能力", PRODUCT_DESC),
        ("导入文件数", len(input_files)),
        ("标准化记录数", len(df)),
        ("资产总数", len(ips)),
        ("涉及风险资产", len(risk_assets)),
        ("高危漏洞记录", int((df["is_high_vulnerability"] == "是").sum()) if not df.empty else 0),
        ("高危端口记录", int((df["is_high_risk_port"] == "是").sum()) if not df.empty else 0),
        ("弱口令记录", int((df["is_weak_password"] == "是").sum()) if not df.empty else 0),
        ("未命中两高一弱", int((df["risk_type"] == "").sum()) if not df.empty else 0),
    ]
    if not df.empty:
        for platform, count in df["source_platform"].value_counts().items():
            items.append((f"来源平台：{platform}", int(count)))
    return _neutralize_formulas(pd.DataFrame(items, columns=["指标", "数量/信息"]))


def export_excel(df: pd.DataFrame, output_path: str, input_files: List[str], mask_passwords: bool = True) -> None:
    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)

    hv = df[df["is_high_vulnerability"] == "是"].copy()
    hp = df[df["is_high_risk_port"] == "是"].copy()
    wp = df[df["is_weak_password"] == "是"].copy()
    unknown = df[df["risk_type"] == ""].copy()

    fd, temp_name = tempfile.mkstemp(prefix=f".{output.stem}-", suffix=".xlsx", dir=output.parent)
    os.close(fd)
    temp_output = Path(temp_name)
    try:
        _write_excel(df, temp_output, input_files, mask_passwords)
        os.replace(temp_output, output)
    finally:
        temp_output.unlink(missing_ok=True)


def _write_excel(df: pd.DataFrame, output: Path, input_files: List[str], mask_passwords: bool) -> None:
    hv = df[df["is_high_vulnerability"] == "是"].copy()
    hp = df[df["is_high_risk_port"] == "是"].copy()
    wp = df[df["is_weak_password"] == "是"].copy()
    unknown = df[df["risk_type"] == ""].copy()

    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        build_summary(df, input_files).to_excel(writer, sheet_name="汇总", index=False)
        _for_export(hv, mask_passwords).to_excel(writer, sheet_name="高危漏洞", index=False)
        _for_export(hp, mask_passwords).to_excel(writer, sheet_name="高危端口", index=False)
        _for_export(wp, mask_passwords).to_excel(writer, sheet_name="弱口令", index=False)
        _for_export(unknown, mask_passwords).to_excel(writer, sheet_name="未识别数据", index=False)
        _for_export(df, mask_passwords).to_excel(writer, sheet_name="全部标准化数据", index=False)

    wb = load_workbook(output)
    header_fill = PatternFill("solid", fgColor="16324F")
    header_font = Font(bold=True, color="EAF4FF")
    alt_fill = PatternFill("solid", fgColor="EEF5FB")
    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for row_idx, row in enumerate(ws.iter_rows(min_row=2), start=2):
            if row_idx % 2 == 0:
                for cell in row:
                    cell.fill = alt_fill
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        for col_idx, col in enumerate(ws.columns, start=1):
            max_len = 0
            for cell in list(col)[:300]:
                value = "" if cell.value is None else str(cell.value)
                max_len = max(max_len, min(len(value), 60))
            ws.column_dimensions[get_column_letter(col_idx)].width = max(10, min(max_len + 2, 40))
    wb.save(output)
